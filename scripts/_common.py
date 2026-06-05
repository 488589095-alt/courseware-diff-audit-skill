# -*- coding: utf-8 -*-
"""
_common.py — 课件对比检查 skill 的共享解析层。

统一三个审计脚本（completeness / structure / style）对 pptx / docx 的访问：
  · 优先读 **实时文件**（.pptx/.docx）—— 信息最全（含 notes、ea 字体、颜色）。
  · 实时文件缺失时回退读 **dump 文本**（_dump_*.txt，由 dump_pptx.py 产出）——
    case 目录里标杆/模版的原始文件常已不在本机，但 dump 一定在。

对外暴露统一的数据结构，下游脚本无需关心数据来自实时文件还是 dump：

  slide = {
    "no": int, "layout": str,
    "text": str,          # 该页所有文字（含表格单元格）拼接，用于内容匹配
    "notes": str,         # 演讲者备注（仅实时读有；dump 回退为 ""）
    "shapes": [ {name,type,pos:[x,y],size:[w,h],is_picture,is_table,
                 text, font:{sz,bold,color,name,ea,latin}, table:[[cell]]} ]
  }
  deck  = {"file": str, "total": int, "slides": [slide], "source": "live"|"dump"}

  docx  = {"file": str, "source": "live"|"dump",
           "items": [ {idx:int, style:str, bold, sz, name, text:str} ],
           "tables": [ {idx:int, rows:[[cell]]} ] }
"""
import re
import unicodedata
from pathlib import Path

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ── 文本归一化与相似度（内容匹配的核心） ──────────────────────────────────────

_PUNCT = r"[，。、；：？！…—·\-_~（）()\[\]【】「」『』《》<>＜＞“”\"'’‘,.;:?!/\\|*•◆■●→←↑↓①②③④⑤⑥⑦⑧⑨⑩]+"


def norm(s: str) -> str:
    """归一化用于子串/相等比较：NFKC(全角→半角) + 去空白 + 去标点 + 小写。"""
    s = unicodedata.normalize("NFKC", s or "")
    s = re.sub(r"\s+", "", s)
    s = re.sub(_PUNCT, "", s)
    return s.lower()


def tokens(s: str) -> set:
    """切 token 用于重叠度：CJK 按单字 + 拉丁/数字按词。"""
    s = unicodedata.normalize("NFKC", s or "").lower()
    cjk = re.findall(r"[一-鿿]", s)
    words = re.findall(r"[a-z0-9]+", s)
    return set(cjk) | set(words)


def similarity(a: str, b: str) -> float:
    """
    讲义条目 a 对候选文本 b 的相似度 ∈ [0,1]。
    取「子串包含」与「token 重叠（以 a 为分母，衡量 a 被 b 覆盖的比例）」的较大值，
    更贴合「讲义这条内容是否被课件这页覆盖」的语义。
    """
    na, nb = norm(a), norm(b)
    if not na:
        return 0.0
    if na in nb:
        return 1.0
    ta, tb = tokens(a), tokens(b)
    if not ta:
        return 0.0
    overlap = len(ta & tb) / len(ta)        # a 的 token 有多少被 b 覆盖
    # 子串部分覆盖加成：a 的较长连续片段是否出现在 b
    longest = _longest_common_substr_ratio(na, nb)
    return max(overlap, longest)


def _longest_common_substr_ratio(a: str, b: str) -> float:
    """最长公共子串长度 / len(a)。短串快速 DP。"""
    if not a or not b:
        return 0.0
    # 限制规模避免超长文本 O(n*m) 爆炸
    a, b = a[:2000], b[:4000]
    prev = [0] * (len(b) + 1)
    best = 0
    for i in range(1, len(a) + 1):
        cur = [0] * (len(b) + 1)
        ai = a[i - 1]
        for j in range(1, len(b) + 1):
            if ai == b[j - 1]:
                cur[j] = prev[j - 1] + 1
                if cur[j] > best:
                    best = cur[j]
        prev = cur
    return best / len(a)


# ── 实时读 pptx ────────────────────────────────────────────────────────────────

def _emu_in(v):
    try:
        from pptx.util import Emu
        return round(Emu(v).inches, 2)
    except Exception:
        return None


def _run_props(r):
    p = {"sz": None, "bold": None, "color": None, "name": None, "ea": None, "latin": None}
    try:
        f = r.font
        if f.size:
            p["sz"] = round(f.size.pt, 1)
        p["bold"] = f.bold
        try:
            if f.color is not None and f.color.type is not None:
                p["color"] = str(f.color.rgb)
        except Exception:
            pass
        p["name"] = f.name
        rpr = r._r.find(f"{A}rPr")
        if rpr is not None:
            for tag in ("latin", "ea"):
                e = rpr.find(f"{A}{tag}")
                if e is not None:
                    p[tag] = e.get("typeface")
    except Exception:
        pass
    return p


def _shape_font(sh):
    """取该 shape 第一个非空 run 的字体属性作为代表。"""
    if not sh.has_text_frame:
        return None
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            if r.text and r.text.strip():
                return _run_props(r)
    return None


def _shape_dict_live(sh):
    try:
        st = str(sh.shape_type)
    except Exception:
        st = "?"
    d = {"name": getattr(sh, "name", ""), "type": st,
         "pos": None, "size": None, "is_picture": False, "is_table": False,
         "text": "", "font": None, "table": None}
    try:
        if sh.left is not None:
            d["pos"] = [_emu_in(sh.left), _emu_in(sh.top)]
            d["size"] = [_emu_in(sh.width), _emu_in(sh.height)]
    except Exception:
        pass
    if st != "?" and "PICTURE" in st:
        d["is_picture"] = True
    try:
        if sh.has_table:
            d["is_table"] = True
            rows = []
            for row in sh.table.rows:
                rows.append([c.text.replace("\n", " ").strip() for c in row.cells])
            d["table"] = rows
    except Exception:
        pass
    try:
        if sh.has_text_frame and sh.text_frame.text.strip():
            d["text"] = " / ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
            d["font"] = _shape_font(sh)
    except Exception:
        pass
    return d


def load_pptx_live(path):
    from pptx import Presentation
    prs = Presentation(str(path))
    slides = []
    for i, s in enumerate(prs.slides, 1):
        try:
            layout = s.slide_layout.name
        except Exception:
            layout = "?"
        shapes, texts = [], []
        for sh in s.shapes:
            d = _shape_dict_live(sh)
            shapes.append(d)
            if d["text"]:
                texts.append(d["text"])
            if d["table"]:
                for row in d["table"]:
                    texts.append(" ".join(c for c in row if c))
        notes = ""
        try:
            if s.has_notes_slide:
                notes = s.notes_slide.notes_text_frame.text or ""
        except Exception:
            pass
        slides.append({"no": i, "layout": layout,
                       "text": "  ".join(texts), "notes": notes, "shapes": shapes})
    return {"file": Path(path).name, "total": len(slides), "slides": slides, "source": "live"}


# ── 回退：解析 pptx dump 文本（dump_pptx.py 的格式） ────────────────────────────
#
# # <path>
# slide_size = W x H in
# total_slides = N
#
# ===== SLIDE 1 =====
# [layout: <name>]
#   <TYPE> name='..' pos=(x,y) size=(wxh) [PICTURE|TABLE]
#       TEXT: a / b / c
#       FONT: sz=18.0,bold=True,col=FF0000,name=宋体
#       TABLE 3x4:
#         | c | c | c | c

_SHAPE_RE = re.compile(r"^\s{2}<(?P<type>[^>]*)>\s+name='(?P<name>[^']*)'\s+pos=\((?P<pos>[^)]*)\)\s+size=\((?P<size>[^)]*)\)\s*(?P<tag>\[PICTURE\]|\[TABLE\])?")
_FONT_RE = re.compile(r"sz=(?P<sz>[^,]*),bold=(?P<bold>[^,]*),col=(?P<col>[^,]*),name=(?P<name>.*)$")


def _parse_pair(s):
    out = []
    for part in s.split(","):
        part = part.strip().rstrip("in").strip()
        try:
            out.append(float(part))
        except Exception:
            out.append(None)
    return out[:2] if len(out) >= 2 else None


def load_pptx_dump(path):
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    slides = []
    cur = None      # current slide dict
    cur_sh = None   # current shape dict
    in_table = False
    for line in text.splitlines():
        m = re.match(r"^===== SLIDE (\d+) =====", line)
        if m:
            if cur:
                slides.append(cur)
            cur = {"no": int(m.group(1)), "layout": "?", "text": "", "notes": "", "shapes": [], "_texts": []}
            cur_sh = None
            in_table = False
            continue
        if cur is None:
            continue
        lm = re.match(r"^\[layout:\s*(.*)\]\s*$", line)
        if lm:
            cur["layout"] = lm.group(1).strip()
            continue
        sm = _SHAPE_RE.match(line)
        if sm:
            tag = sm.group("tag") or ""
            cur_sh = {"name": sm.group("name"), "type": sm.group("type"),
                      "pos": _parse_pair(sm.group("pos")), "size": _parse_pair(sm.group("size")),
                      "is_picture": "PICTURE" in tag, "is_table": "TABLE" in tag,
                      "text": "", "font": None, "table": [] if "TABLE" in tag else None}
            cur["shapes"].append(cur_sh)
            in_table = False
            continue
        if line.startswith("      TEXT: ") and cur_sh is not None:
            cur_sh["text"] = line[len("      TEXT: "):]
            cur["_texts"].append(cur_sh["text"])
            continue
        if line.startswith("      FONT: ") and cur_sh is not None:
            fm = _FONT_RE.search(line[len("      FONT: "):])
            if fm:
                def _v(x):
                    x = x.strip()
                    return None if x in ("None", "") else x
                sz = _v(fm.group("sz"))
                cur_sh["font"] = {
                    "sz": float(sz) if sz and re.match(r"^[\d.]+$", sz) else None,
                    "bold": {"True": True, "False": False}.get(_v(fm.group("bold") or "")),
                    "color": _v(fm.group("col")),
                    "name": _v(fm.group("name")),
                    "ea": None, "latin": None,
                }
            continue
        if re.match(r"^      TABLE \d+x\d+:", line) and cur_sh is not None:
            in_table = True
            if cur_sh.get("table") is None:
                cur_sh["table"] = []
            cur_sh["is_table"] = True
            continue
        if in_table and line.startswith("        | "):
            cells = [c.strip() for c in line[len("        | "):].split(" | ")]
            cur_sh["table"].append(cells)
            cur["_texts"].append(" ".join(c for c in cells if c))
            continue
    if cur:
        slides.append(cur)
    for s in slides:
        s["text"] = "  ".join(s.pop("_texts"))
    return {"file": Path(path).name, "total": len(slides), "slides": slides, "source": "dump"}


def load_deck(path):
    """统一入口：.pptx→实时读；.txt→dump 回退读。"""
    p = Path(path)
    if p.suffix.lower() == ".pptx":
        return load_pptx_live(p)
    return load_pptx_dump(p)


# ── docx：实时读 + dump 回退 ────────────────────────────────────────────────────

def load_docx_live(path):
    import docx
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(str(path))
    items, tables = [], []
    pi = ti = 0
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            p = Paragraph(child, doc)
            txt = p.text
            if txt.strip():
                bold = sz = name = None
                for r in p.runs:
                    if r.text.strip():
                        bold = r.bold
                        sz = r.font.size.pt if r.font.size else None
                        name = r.font.name
                        break
                items.append({"idx": pi, "style": p.style.name if p.style else "",
                              "bold": bold, "sz": sz, "name": name, "text": txt})
            pi += 1
        elif isinstance(child, CT_Tbl):
            t = Table(child, doc)
            rows = [[c.paragraphs[0].text if c.paragraphs else "" for c in row.cells]
                    for row in t.rows]
            tables.append({"idx": ti, "rows": rows})
            ti += 1
    return {"file": Path(path).name, "source": "live", "items": items, "tables": tables}


def load_docx_dump(path):
    """
    解析 docx dump，兼容两种格式（不同 case 的 extractor 产出不一）：
      格式 A (dump_pptx.py):  `P{i} <style> [meta]: text`  + `--- TABLE k ---`
      格式 B (古诗等 extractor): `[P|style] text`            + `[TABLE rxc]`
    两种格式的表格行均为 `  | a | b | c`。
    """
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    items, tables = [], []
    cur_tbl = None
    pidx = 0
    for line in text.splitlines():
        # 格式 A 段落
        pm = re.match(r"^P(\d+)\s+<([^>]*)>\s+(\[[^\]]*\])?\s*:\s*(.*)$", line)
        if pm:
            cur_tbl = None
            meta = pm.group(3) or ""
            bold = sz = name = None
            bm = re.search(r"b=([^,]*)", meta)
            sm = re.search(r"sz=([^,]*)", meta)
            nm = re.search(r"name=([^\]]*)", meta)
            if bm:
                bold = {"True": True, "False": False}.get(bm.group(1).strip())
            if sm and re.match(r"^[\d.]+$", sm.group(1).strip()):
                sz = float(sm.group(1).strip())
            if nm and nm.group(1).strip() not in ("None", ""):
                name = nm.group(1).strip()
            pidx = int(pm.group(1))
            items.append({"idx": pidx, "style": pm.group(2).strip(),
                          "bold": bold, "sz": sz, "name": name, "text": pm.group(4)})
            continue
        # 格式 B 段落: [P|Style] text
        pb = re.match(r"^\[P\|([^\]]*)\]\s?(.*)$", line)
        if pb:
            cur_tbl = None
            items.append({"idx": pidx, "style": pb.group(1).strip(),
                          "bold": None, "sz": None, "name": None, "text": pb.group(2)})
            pidx += 1
            continue
        # 表格头：格式 A `--- TABLE k (rxc) ---` 或格式 B `[TABLE rxc]`
        tm = re.match(r"^--- TABLE (\d+) \((\d+)x(\d+)\) ---", line)
        tb = re.match(r"^\[TABLE (\d+)x(\d+)\]", line)
        if tm:
            cur_tbl = {"idx": int(tm.group(1)), "rows": []}
            tables.append(cur_tbl)
            continue
        if tb:
            cur_tbl = {"idx": len(tables), "rows": []}
            tables.append(cur_tbl)
            continue
        if cur_tbl is not None and line.startswith("  | "):
            cur_tbl["rows"].append([c.strip() for c in line[len("  | "):].split(" | ")])
            continue
    return {"file": Path(path).name, "source": "dump", "items": items, "tables": tables}


def load_handout(path):
    """统一入口：.docx→实时读；.txt→dump 回退读。"""
    p = Path(path)
    if p.suffix.lower() == ".docx":
        return load_docx_live(p)
    return load_docx_dump(p)
