# -*- coding: utf-8 -*-
"""
style_fidelity.py — 课件 vs 模版+标杆 的样式保真度（数据驱动，不硬编码字体）。

期望基准两路：
  · 模版主题（template_spec.json）：合法字体集（major/minor）、主题配色 → 课件该继承。
  · 标杆调色板：老师精备的「期望视觉」→ 颜色/字号/字体分布，课件该贴近。

产出：
  · hard_rules     通用硬规则（跨学科）：WPS 冻结 / 行距 / **中文回退等线**（ea=等线/DengXian）
  · baseline       模版主题字体 + 配色（来自 template_spec.json，缺则标注 limited）
  · palette        课件 / 标杆 各自的 颜色·字体·字号 分布（课件实时读=run级；标杆dump=shape级）
  · palette_diff   标杆用而课件没用的颜色/字号/字体（=样式没贴上），及反向
  · style_calibration_hints  机械可比的几个关键元素差异（答案红、分隔标题色号等）的线索

style_calibration_table（语义元素逐行）由 Claude 用本文件 palette 事实 + 截图填写。

用法:
  python3 style_fidelity.py --courseware <课件.pptx> \
      [--benchmark <标杆.pptx|dump.txt>] [--template-spec template_spec.json] -o style_diff.json
"""
import argparse
import json
import re
from collections import Counter
from pathlib import Path

import _common as C

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
FALLBACK_EA = re.compile(r"等线|dengxian|deng\s*xian", re.I)
RED = re.compile(r"^(FF0000|FF5050|E_?A4D54|7D292D|CC0000|FF3333|D9001B|C00000)$", re.I)


# ── 调色板提取 ──────────────────────────────────────────────────────────────────

def palette_from_pptx(path):
    """实时 pptx：run 级，最全（含 ea/latin/color/size）。"""
    from pptx import Presentation
    prs = Presentation(str(path))
    colors, fonts_ea, fonts_latin, fonts_name, sizes = (Counter() for _ in range(5))
    fallback_pages, cjk_no_ea_pages = set(), set()
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not (r.text and r.text.strip()):
                        continue
                    pr = C._run_props(r)
                    has_cjk = bool(re.search(r"[一-鿿]", r.text))
                    if pr["color"]:
                        colors[pr["color"]] += 1
                    if pr["sz"]:
                        sizes[pr["sz"]] += 1
                    if pr["ea"]:
                        fonts_ea[pr["ea"]] += 1
                        # 仅当该 run 含中文时，ea=等线 才真的会显示为等线（纯英文走 latin）
                        if has_cjk and FALLBACK_EA.search(pr["ea"]):
                            fallback_pages.add(i)
                    if pr["latin"]:
                        fonts_latin[pr["latin"]] += 1
                    if pr["name"]:
                        fonts_name[pr["name"]] += 1
                    # 含中文但未显式设 ea → 可能回退主题默认（等线）
                    if not pr["ea"] and has_cjk:
                        cjk_no_ea_pages.add(i)
    return {"colors": colors, "fonts_ea": fonts_ea, "fonts_latin": fonts_latin,
            "fonts_name": fonts_name, "sizes": sizes,
            "_fallback_pages": sorted(fallback_pages),
            "_cjk_no_ea_pages": sorted(cjk_no_ea_pages)}


def palette_from_deck(deck):
    """dump 回退：shape 级（每 shape 取首个 run），用于标杆。"""
    colors, fonts_name, sizes = Counter(), Counter(), Counter()
    for s in deck["slides"]:
        for sh in s["shapes"]:
            f = sh.get("font")
            if not f:
                continue
            if f.get("color"):
                colors[f["color"]] += 1
            if f.get("sz"):
                sizes[f["sz"]] += 1
            if f.get("name"):
                fonts_name[f["name"]] += 1
    return {"colors": colors, "fonts_ea": Counter(), "fonts_latin": Counter(),
            "fonts_name": fonts_name, "sizes": sizes,
            "_fallback_pages": [], "_cjk_no_ea_pages": []}


def get_palette(path):
    p = Path(path)
    if p.suffix.lower() == ".pptx":
        return palette_from_pptx(p)
    return palette_from_deck(C.load_deck(p))


def _top(counter, n=12):
    return [{"value": str(k), "count": v} for k, v in counter.most_common(n)]


# ── 硬规则（跨学科通用，复用 structural_checks 的 WPS / 行距） ──────────────────

def hard_rules(courseware_pptx, fallback_pages, cjk_no_ea_pages):
    out = []
    try:
        import structural_checks as SC
        from pptx import Presentation
        prs = Presentation(str(courseware_pptx))
        for rule, status, detail in SC.check_wps_patterns(prs, Path(courseware_pptx).name):
            out.append({"rule": rule, "status": status, "detail": detail})
        for rule, status, detail in SC.check_line_spacing(prs, Path(courseware_pptx).name):
            out.append({"rule": rule, "status": status, "detail": detail})
    except Exception as e:
        out.append({"rule": "结构检查", "status": "WARN", "detail": f"无法运行: {e}"})
    # 通用字体回退检查（替代 check_fonts 的微软雅黑硬编码）
    if fallback_pages:
        out.append({"rule": "字体:中文回退等线", "status": "FAIL",
                    "pages": fallback_pages,
                    "detail": f"{len(fallback_pages)}页 ea字体=等线/DengXian（应显式设中文字体，否则WPS/PPT显示等线）"})
    else:
        out.append({"rule": "字体:中文回退等线", "status": "PASS", "detail": "无 ea=等线"})
    if cjk_no_ea_pages:
        out.append({"rule": "字体:中文未显式设ea", "status": "WARN",
                    "pages": cjk_no_ea_pages[:30],
                    "detail": f"{len(cjk_no_ea_pages)}页 有中文run未显式设a:ea（依赖主题默认,有回退等线风险）"})
    return out


# ── 模版基准 ────────────────────────────────────────────────────────────────────

def load_baseline(spec_path):
    if spec_path and Path(spec_path).exists():
        spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
        th = spec.get("theme", {})
        return {"source": "template_spec.json", "major_font": th.get("major_font"),
                "minor_font": th.get("minor_font"), "theme_colors": th.get("colors", {})}
    return {"source": "none", "note": "无 template_spec.json（建议先 dissect_template 模版）；"
            "样式期望退化为仅以标杆调色板为准"}


def run(courseware_path, benchmark_path=None, template_spec=None):
    cw_pal = get_palette(courseware_path)
    bm_pal = get_palette(benchmark_path) if benchmark_path else None
    baseline = load_baseline(template_spec)
    hr = hard_rules(courseware_path, cw_pal["_fallback_pages"], cw_pal["_cjk_no_ea_pages"]) \
        if Path(courseware_path).suffix.lower() == ".pptx" else \
        [{"rule": "硬规则", "status": "WARN", "detail": "课件非实时pptx，跳过硬规则"}]

    def export(pal):
        return {"colors": _top(pal["colors"]), "fonts_ea": _top(pal["fonts_ea"]),
                "fonts_latin": _top(pal["fonts_latin"]), "fonts_name": _top(pal["fonts_name"]),
                "sizes": _top(pal["sizes"])}

    palette = {"courseware": export(cw_pal)}
    palette_diff = None
    hints = []
    if bm_pal:
        palette["benchmark"] = export(bm_pal)
        cwc, bmc = set(cw_pal["colors"]), set(bm_pal["colors"])
        cws, bms = set(cw_pal["sizes"]), set(bm_pal["sizes"])
        cwf = set(cw_pal["fonts_name"]) | set(cw_pal["fonts_ea"])
        bmf = set(bm_pal["fonts_name"]) | set(bm_pal["fonts_ea"])
        palette_diff = {
            "colors_benchmark_only": sorted(bmc - cwc),
            "colors_courseware_only": sorted(cwc - bmc),
            "sizes_benchmark_only": sorted(bms - cws),
            "fonts_benchmark_only": sorted(bmf - cwf),
            "fonts_courseware_only": sorted(cwf - bmf),
        }
        # 关键元素线索：答案红是否用上
        bm_red = sorted(c for c in bmc if RED.match(c))
        cw_red = sorted(c for c in cwc if RED.match(c))
        if bm_red and not cw_red:
            hints.append({"element": "强调红(答案/标签)", "issue": "标杆用红强调而课件未见红色字",
                          "benchmark_red": bm_red, "courseware_red": cw_red})
        elif bm_red and cw_red and set(bm_red) != set(cw_red):
            hints.append({"element": "强调红色值", "issue": "红色值与标杆不一致",
                          "benchmark_red": bm_red, "courseware_red": cw_red})

    return {
        "courseware": Path(courseware_path).name,
        "benchmark": Path(benchmark_path).name if benchmark_path else None,
        "baseline": baseline,
        "hard_rules": hr,
        "palette": palette,
        "palette_diff": palette_diff,
        "style_calibration_hints": hints,
        "note": "style_calibration_table（题/诗/答案/分隔标题等语义元素逐行）由 Claude 用本"
                "palette + 截图填写；脚本只给颜色/字体/字号分布事实",
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--courseware", required=True)
    ap.add_argument("--benchmark", default=None)
    ap.add_argument("--template-spec", default=None)
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()
    res = run(a.courseware, a.benchmark, a.template_spec)
    Path(a.out).write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    print("样式保真:")
    for h in res["hard_rules"]:
        mark = {"PASS": "✓", "FAIL": "✗", "WARN": "△"}.get(h["status"], "·")
        print(f"  {mark} {h['rule']}: {h['detail']}")
    if res["palette_diff"]:
        pd = res["palette_diff"]
        print(f"  标杆有/课件无 颜色: {pd['colors_benchmark_only'][:8]}")
        print(f"  标杆有/课件无 字号: {pd['sizes_benchmark_only'][:8]}")
        print(f"  标杆有/课件无 字体: {pd['fonts_benchmark_only'][:8]}")
    for h in res["style_calibration_hints"]:
        print(f"  ⚑ {h['element']}: {h['issue']}")
    print(f"✅ {a.out}")


if __name__ == "__main__":
    main()
